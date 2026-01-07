from pptx import Presentation
from pptx.util import Inches
from pathlib import Path

OUT = Path('d:/moqt-project/outputs')
prs = Presentation()
blank = prs.slide_layouts[6]

svgs = [OUT / 'chart_given_large.svg', OUT / 'chart_sorted_large.svg']
for svg in svgs:
    slide = prs.slides.add_slide(blank)
    left = Inches(0.5)
    top = Inches(1)
    width = Inches(9)
    try:
        pic = slide.shapes.add_picture(str(svg), left, top, width=width)
        print(f'Embedded {svg} as picture')
    except Exception as e:
        print(f'Failed to embed {svg}: {e}')
        # fallback: add text box with file path
        tx = slide.shapes.add_textbox(left, top, width, Inches(1))
        tx.text = f'Image unavailable in PPTX. See file: {svg}'

pptx_path = OUT / 'charts_presentation.pptx'
prs.save(pptx_path)
print('Saved PPTX:', pptx_path)
